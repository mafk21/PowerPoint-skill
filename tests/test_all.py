#!/usr/bin/env python3
"""Repo test suite (stdlib unittest — no pytest needed).

Covers: builders registered, Morph/timings XML, Chart Lab anchors, icon
engine, RTL polish, bilingual guards, Arabic-Indic numerals, Hijri math,
Palette Lab contrast floors, theme merge, DB census, QA gate both ways,
draft heuristics, portrait formats, overview data lint.

Run:  python -m unittest discover -s tests -v
"""
import importlib.util
import json
import os
import re
import sqlite3
import sys
import tempfile
import unittest
from zipfile import ZipFile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "generator"))
sys.path.insert(0, os.path.join(ROOT, "tools"))

import generate_pptx as G          # noqa: E402
import portrait                    # noqa: E402
import palette_lab                 # noqa: E402
import draft                       # noqa: E402


def load_tool(name):
    spec = importlib.util.spec_from_file_location(
        name, os.path.join(ROOT, "tools", name + ".py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def slide_xmls(path):
    z = ZipFile(path)
    names = sorted((n for n in z.namelist()
                    if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                   key=lambda n: int(re.search(r"(\d+)", n).group(1)))
    return [z.read(n).decode() for n in names]


TMP = tempfile.mkdtemp(prefix="ppt-tests-")


def build(spec, name):
    out = os.path.join(TMP, name)
    G.build(spec, out)
    return out


BASE = {"theme": "teal-orbit", "brand": "T", "slides": [
    {"type": "title", "title_lines": ["T"], "subtitle": "s", "notes": "hello world this is a script"},
    {"type": "hero_stat", "stat": "87%", "label": "renewed", "notes": "n"},
    {"type": "closing", "title_lines": ["Go"], "notes": "n"}]}


class Generator(unittest.TestCase):
    def test_builder_registry(self):
        self.assertGreaterEqual(len(G.BUILDERS), 22)
        for b in ("line_chart", "donut", "waterfall", "matrix", "funnel",
                  "bilingual"):
            self.assertIn(b, G.BUILDERS)
        self.assertEqual(len(portrait.BUILDERS), 8)

    def test_morph_and_no_morph_on_first(self):
        xmls = slide_xmls(build(BASE, "a.pptx"))
        self.assertNotIn("p14:morph", xmls[0])
        self.assertTrue(all("p14:morph" in x for x in xmls[1:]))

    def test_auto_timings(self):
        spec = dict(BASE, timings=True)
        xmls = slide_xmls(build(spec, "b.pptx"))
        self.assertIn("p:fade", xmls[0])          # slide 1 carries timing
        for x in xmls:
            self.assertRegex(x, r'advTm="\d+"')

    def test_waterfall_anchors_and_vermilion(self):
        spec = {"theme": "teal-orbit", "slides": [
            {"type": "title", "title_lines": ["T"]},
            {"type": "waterfall", "title": "Bridge from A to B",
             "steps": [{"label": "Start", "value": 100, "kind": "start"},
                       {"label": "Gain", "value": 20},
                       {"label": "Loss", "value": -35},
                       {"label": "End", "value": 85, "kind": "total"}]}]}
        x = slide_xmls(build(spec, "c.pptx"))[1]
        self.assertIn("!!wf-0", x)
        self.assertIn("D55E00", x)                # Okabe vermilion down-step

    def test_icon_engine(self):
        self.assertEqual(G._icon_glyph({"icon": "growth"}), "↗")
        self.assertEqual(G._icon_glyph({"glyph": "★", "icon": "x"}), "★")
        self.assertEqual(G._icon_glyph({}), "◆")

    def test_external_themes_merged(self):
        self.assertIn("saudi-green", G.THEMES)    # via themes.json
        self.assertIn("coral-pop", G.THEMES)
        self.assertGreaterEqual(len(G.THEMES), 19)


class Arabic(unittest.TestCase):
    def test_rtl_polish(self):
        spec = {"theme": "emerald-ring", "rtl": True, "slides": [
            {"type": "title", "title_lines": ["منصة"], "subtitle": "نص"},
            {"type": "process", "title": "كيف يعمل",
             "steps": [{"head": "أ", "body": "١"}, {"head": "ب", "body": "٢"},
                       {"head": "ج", "body": "٣"}]}]}
        xmls = slide_xmls(build(spec, "ar.pptx"))
        self.assertIn('rtl="1"', xmls[0])
        self.assertIn("flipH", xmls[1])           # chevrons flipped
        self.assertIn("Cairo", xmls[0])

    def test_bilingual_guards_and_indic(self):
        spec = {"theme": "saudi-green", "rtl": True,
                "numerals": "arabic-indic", "slides": [
                    {"type": "title", "title_lines": ["منصة"]},
                    {"type": "bilingual", "ar_title": "عنوان 42",
                     "en_title": "Title 42",
                     "ar_bullets": [{"head": "نقطة 3", "body": "شرح"}],
                     "en_bullets": [{"head": "Point 3", "body": "Text"}]}]}
        x = slide_xmls(build(spec, "bi.pptx"))[1]
        en = [p for p in re.findall(r"<a:p>.*?</a:p>", x, re.S)
              if "Point" in p]
        self.assertTrue(en and 'rtl="1"' not in en[0])
        self.assertIn("٤٢", x)                    # Arabic title digits converted
        self.assertIn("Title 42", x)              # English digits untouched

    def test_hijri_math(self):
        self.assertEqual(G.gregorian_to_hijri(2026, 10, 19), (1448, 5, 7))
        self.assertEqual(G.gregorian_to_hijri(622, 7, 19), (1, 1, 1))
        out = G._expand_tokens({"meta": "{hijri}"}, {"date": "2026-10-19"})
        self.assertIn("1448هـ", out["meta"])


class ColorEngine(unittest.TestCase):
    def test_palette_lab_floors(self):
        for seed, dark in (("FF6B4A", False), ("0AA9A0", True),
                           ("7C3AED", False), ("006C35", True),
                           ("D4A537", True), ("0284C7", False)):
            theme, report = palette_lab.build_theme("t", seed, None, dark)
            for line, ok in report:
                self.assertTrue(ok, f"{seed} dark={dark}: {line}")

    def test_all_themes_pass_mutued_floor(self):
        for name, t in G.THEMES.items():
            cm = palette_lab.contrast(t["muted"], t["bg"])
            self.assertGreaterEqual(cm, 4.5, f"{name} muted:bg={cm:.2f}")


class Database(unittest.TestCase):
    def test_census(self):
        db = os.path.join(ROOT, "database", "powerpoint.db")
        self.assertTrue(os.path.exists(db))
        c = sqlite3.connect(db)
        self.assertEqual(c.execute("SELECT COUNT(*) FROM creators"
                                   ).fetchone()[0], 36)
        self.assertGreaterEqual(
            c.execute("SELECT COUNT(*) FROM palettes").fetchone()[0], 21)
        names = {r[0] for r in c.execute("SELECT name FROM structures")}
        self.assertTrue(any("pitch" in n.lower() for n in names))
        # FTS works
        hits = c.execute("SELECT COUNT(*) FROM search_index "
                         "WHERE search_index MATCH 'morph'").fetchone()[0]
        self.assertGreater(hits, 0)


class QA(unittest.TestCase):
    def test_green_deck_passes(self):
        qa = load_tool("qa_deck")
        self.assertEqual(qa.main([build(BASE, "qa-ok.pptx")]), 0)

    def test_tiny_text_fails(self):
        from pptx import Presentation
        from pptx.util import Inches, Pt
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        r = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4),
                                 Inches(1)).text_frame.paragraphs[0].add_run()
        r.text = "tiny"
        r.font.size = Pt(6)
        bad = os.path.join(TMP, "qa-bad.pptx")
        prs.save(bad)
        qa = load_tool("qa_deck")
        self.assertEqual(qa.main([bad]), 1)


class CreatorTools(unittest.TestCase):
    def test_draft_classification(self):
        notes = ("# Acme\nWe fix widgets\n\n# Data\n"
                 "Jan: 3\nFeb: 5\nMar: 9\n\n42%\nof users stay\n\n"
                 "# Ask\nWe are raising $1M. Contact us and join us.\n")
        spec, report = draft.draft_spec(notes)
        types = [s["type"] for s in spec["slides"]]
        self.assertEqual(types, ["title", "chart", "hero_stat", "closing"])
        self.assertTrue(all(s.get("notes") for s in spec["slides"]))

    def test_portrait_formats(self):
        for fmt, (w, h) in portrait.FORMATS.items():
            out = os.path.join(TMP, f"{fmt}.pptx")
            portrait.build_portrait(
                {"format": fmt, "theme": "teal-orbit", "brand": "T",
                 "slides": [{"type": "hook", "title_lines": ["Hi"],
                             "notes": "vo line here please ok"},
                            {"type": "bigstat", "stat": "3x",
                             "label": "proof"}]}, out)
            from pptx import Presentation
            prs = Presentation(out)
            self.assertAlmostEqual(prs.slide_width.inches, w, places=2)
            self.assertAlmostEqual(prs.slide_height.inches, h, places=2)
            if fmt == "reel":
                self.assertRegex(slide_xmls(out)[0], r'advTm="\d+"')

    def test_db_lint_clean(self):
        lint = load_tool("db_lint")
        self.assertEqual(lint.main([]), 0)

    def test_examples_qa_green(self):
        qa = load_tool("qa_deck")
        ex = os.path.join(ROOT, "examples")
        for fn in sorted(os.listdir(ex)):
            if fn.endswith(".pptx"):
                self.assertEqual(qa.main([os.path.join(ex, fn)]), 0, fn)


class Showpiece(unittest.TestCase):
    def test_htmldeck(self):
        hd = load_tool("htmldeck")
        spec_path = os.path.join(TMP, "hd.json")
        json.dump(BASE, open(spec_path, "w"))
        out = os.path.join(TMP, "hd.html")
        self.assertEqual(hd.main(["htmldeck", spec_path, "--out", out]), 0)
        doc = open(out, encoding="utf-8").read()
        self.assertEqual(doc.count('<section class="slide"'),
                         len(BASE["slides"]))
        self.assertIn("--accent:", doc)

    def test_site_generates(self):
        site = load_tool("site")
        outdir = os.path.join(TMP, "site")
        self.assertEqual(site.main(["site", "--outdir", outdir]), 0)
        doc = open(os.path.join(outdir, "index.html"), encoding="utf-8").read()
        self.assertIn("creators", doc)          # census table present
        self.assertIn("saudi-green", doc)       # theme gallery live
        self.assertTrue(os.path.exists(os.path.join(outdir, "demo.html")))


if __name__ == "__main__":
    unittest.main(verbosity=2)
