#!/usr/bin/env python3
"""Deck Anatomy — reverse-engineer any .pptx: palette, fonts, Morph map,
per-slide content census, and a reusable generator theme guessed from it.

Usage:
  python tools/anatomy.py deck.pptx
  python tools/anatomy.py deck.pptx --write NAME     # also save guessed theme
  python tools/anatomy.py deck.pptx --out report.md
"""
import argparse
import os
import re
import sys
from collections import Counter
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from palette_lab import build_theme, contrast, hsl, save_theme, show  # noqa: E402


def morph_map(path, n):
    z = ZipFile(path)
    out = []
    for i in range(1, n + 1):
        try:
            x = z.read(f"ppt/slides/slide{i}.xml").decode("utf-8", "ignore")
        except KeyError:
            out.append("?")
            continue
        if "p14:morph" in x:
            out.append("M")
        elif "<p:transition" in x:
            out.append("t")
        else:
            out.append(".")
    return "".join(out)


def solid_hex(shape):
    try:
        if shape.fill and shape.fill.type == MSO_FILL.SOLID:
            return str(shape.fill.fore_color.rgb)
    except Exception:
        pass
    return None


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)
        else:
            yield sh


def analyze(path):
    prs = Presentation(path)
    W, H = prs.slide_width.inches, prs.slide_height.inches
    fill_area, run_color = Counter(), Counter()
    fonts, sizes = Counter(), []
    bg_colors = Counter()
    rows, notes_n = [], 0

    for i, slide in enumerate(prs.slides, 1):
        try:
            if slide.background.fill.type == MSO_FILL.SOLID:
                bg_colors[str(slide.background.fill.fore_color.rgb)] += 1
        except Exception:
            pass
        kinds = set()
        words, max_pt = 0, 0
        for sh in walk(slide.shapes):
            if sh.shape_type == MSO_SHAPE_TYPE.CHART:
                kinds.add("chart")
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                kinds.add("picture")
            elif sh.has_table:
                kinds.add("table")
            hx = solid_hex(sh)
            if hx and sh.width and sh.height:
                fill_area[hx] += sh.width.inches * sh.height.inches
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for run in para.runs:
                        txt = run.text.strip()
                        if not txt:
                            continue
                        words += len(txt.split())
                        f = run.font
                        if f.name:
                            fonts[f.name] += 1
                        if f.size:
                            sizes.append(f.size.pt)
                            max_pt = max(max_pt, f.size.pt)
                        try:
                            if f.color and f.color.type is not None:
                                run_color[str(f.color.rgb)] += len(txt)
                        except Exception:
                            pass
        has_notes = bool(slide.has_notes_slide and
                         slide.notes_slide.notes_text_frame.text.strip())
        notes_n += has_notes
        guess = ("/".join(sorted(kinds)) if kinds else
                 "hero" if max_pt >= 44 else "text")
        rows.append((i, guess, words, max_pt, "yes" if has_notes else "—"))

    sizes.sort()
    med = sizes[len(sizes) // 2] if sizes else 0
    return dict(n=len(rows), W=W, H=H, morph=morph_map(path, len(rows)),
                fill_area=fill_area, run_color=run_color, fonts=fonts,
                size_min=min(sizes) if sizes else 0, size_med=med,
                size_max=max(sizes) if sizes else 0, bg_colors=bg_colors,
                rows=rows, notes_n=notes_n)


def guess_theme(a):
    # bg: explicit slide background, else largest fill area
    bg = (a["bg_colors"].most_common(1) or [(None, 0)])[0][0]
    if not bg and a["fill_area"]:
        bg = a["fill_area"].most_common(1)[0][0]
    dark = False
    if bg:
        _, l, _ = hsl(bg)
        dark = l < 0.35
    # accent: saturated fills by area, BrandKit-style scoring
    total = sum(a["fill_area"].values()) or 1
    scored = []
    for hexs, area in a["fill_area"].items():
        h, l, s = hsl(hexs)
        if s < 0.22 or l <= 0.08 or l >= 0.93 or hexs == bg:
            continue
        window = 1.0 - min(1.0, abs(l - 0.52) / 0.45)
        scored.append((((area / total) ** 0.5) * (s ** 1.5) * (0.25 + window),
                       hexs, h))
    scored.sort(reverse=True)
    if not scored:
        return None, None
    accent, ah = scored[0][1], scored[0][2]
    accent2 = None
    for _, hexs, h in scored[1:]:
        d = min(abs(h - ah) % 1.0, 1 - (abs(h - ah) % 1.0)) * 360
        if 25 <= d <= 200:
            accent2 = hexs
            break
    theme, report = build_theme("guessed", accent, accent2, dark=dark)
    if bg:  # pin to the deck's real background; card = largest near-bg fill
        theme["bg"] = bg
        _, bgl, _ = hsl(bg)
        for hexs, _area in a["fill_area"].most_common():
            if hexs == bg:
                continue
            _, l, _ = hsl(hexs)
            if 0.02 < abs(l - bgl) < 0.28:
                theme["card"] = hexs
                break
        report = [
            (f"ink:bg {contrast(theme['ink'], bg):.1f}:1 (floor 10)",
             contrast(theme["ink"], bg) >= 10),
            (f"muted:bg {contrast(theme['muted'], bg):.1f}:1 (floor 4.5)",
             contrast(theme["muted"], bg) >= 4.5),
            (f"accent:bg {contrast(theme['accent'], bg):.1f}:1 (floor 3)",
             contrast(theme["accent"], bg) >= 3.0)]
    return (theme, report), bg


def render(path, a, theme_pack, real_bg):
    L = [f"# Deck anatomy — {os.path.basename(path)}", ""]
    L.append(f"- Canvas: {a['W']:.2f}×{a['H']:.2f} in "
             f"({'16:9' if abs(a['W']/a['H'] - 16/9) < 0.02 else 'non-16:9'})")
    L.append(f"- Slides: {a['n']} · notes on {a['notes_n']}/{a['n']}")
    L.append(f"- Morph map: `{a['morph']}`  (M=Morph, t=other transition, .=none)")
    L.append(f"- Fonts: " + (", ".join(f"{f}×{c}" for f, c in a["fonts"].most_common())
                            or "(theme-default runs)"))
    L.append(f"- Type sizes: min {a['size_min']:.0f}pt · median "
             f"{a['size_med']:.0f}pt · max {a['size_max']:.0f}pt")
    L.append("")
    L.append("## Palette by painted area (fills)")
    total = sum(a["fill_area"].values()) or 1
    for hexs, area in a["fill_area"].most_common(6):
        L.append(f"- `#{hexs}` — {area/total:.0%}")
    L.append("")
    L.append("## Text colors by character count")
    tc = sum(a["run_color"].values()) or 1
    for hexs, c in a["run_color"].most_common(5):
        L.append(f"- `#{hexs}` — {c/tc:.0%}")
    L.append("")
    L.append("## Slides")
    L.append("| # | content guess | words | max type | notes |")
    L.append("|---|---|---|---|---|")
    for i, g, w, mp, nt in a["rows"]:
        L.append(f"| {i} | {g} | {w} | {mp:.0f}pt | {nt} |")
    L.append("")
    if theme_pack:
        theme, report = theme_pack
        L.append("## Guessed generator theme")
        L.append(f"- detected slide background: `#{real_bg}` "
                 f"({'dark' if theme['dark'] else 'light'} mode)")
        L.append(f"- roles: bg `#{theme['bg']}` · card `#{theme['card']}` · "
                 f"ink `#{theme['ink']}` · muted `#{theme['muted']}` · "
                 f"accent `#{theme['accent']}` · accent2 `#{theme['accent2']}` · "
                 f"graybar `#{theme['graybar']}`")
        for line, ok in report:
            L.append(f"- {'PASS' if ok else 'WARN'} {line}")
        L.append("- save it with `--write <name>` to reuse in specs")
    else:
        L.append("## Guessed theme: none (no saturated fills found)")
    return "\n".join(L) + "\n"


def main(argv):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("deck")
    ap.add_argument("--write", metavar="NAME", help="save guessed theme under NAME")
    ap.add_argument("--out", metavar="FILE", help="write report to FILE too")
    a = ap.parse_args(argv[1:])
    if not os.path.exists(a.deck):
        print(f"no such file: {a.deck}")
        return 2
    ana = analyze(a.deck)
    theme_pack, real_bg = guess_theme(ana)
    report = render(a.deck, ana, theme_pack, real_bg)
    print(report)
    if a.out:
        with open(a.out, "w", encoding="utf-8") as f:
            f.write(report)
        print(f"[report written to {a.out}]")
    if a.write and theme_pack:
        theme, rep = theme_pack
        total = save_theme(a.write, theme)
        show(a.write, theme, rep)
        print(f"saved — themes.json now holds {total}; use "
              f'"theme": "{a.write}"')
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
