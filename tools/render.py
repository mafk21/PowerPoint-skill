#!/usr/bin/env python3
"""Contact-sheet renderer — see your deck without PowerPoint.

Reads the real geometry of a .pptx (python-pptx) and emits ONE standalone
HTML file: every slide as a scaled wireframe with true positions, fills,
type sizes and text. It is a *proof sheet*, not a pixel clone — charts,
tables and pictures become labeled blocks (chart data is summarized when
readable). Perfect for catching overflow/crowding the QA sentinel flags,
sharing a preview, or diffing two versions visually.

Usage:
  python tools/render.py deck.pptx [--out contact.html] [--width 640]
"""
import argparse
import html
import os
import re
import sys
import zipfile
from collections import Counter

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from anatomy import morph_map, walk  # noqa: E402


def solid(shape):
    try:
        if shape.fill and shape.fill.type == MSO_FILL.SOLID:
            return "#" + str(shape.fill.fore_color.rgb)
    except Exception:
        pass
    return None


def slide_bg(slide, shapes):
    try:
        if slide.background.fill.type == MSO_FILL.SOLID:
            return "#" + str(slide.background.fill.fore_color.rgb)
    except Exception:
        pass
    areas = Counter()
    for sh in shapes:
        c = solid(sh)
        if c and sh.width and sh.height:
            areas[c] += sh.width.inches * sh.height.inches
    return areas.most_common(1)[0][0] if areas else "#FFFFFF"


def radius_for(shape):
    try:
        name = str(shape.auto_type)
    except Exception:
        return "0"
    if "OVAL" in name:
        return "50%"
    if "ROUNDED" in name:
        return "10px"
    return "0"


def text_html(shape, scale):
    tf = shape.text_frame
    paras = []
    for para in tf.paragraphs:
        if not para.runs:
            continue
        align = {PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right"}.get(
            para.alignment, "left")
        rtl = bool(para._pPr is not None and para._pPr.get("rtl") == "1")
        runs = []
        for r in para.runs:
            if not r.text:
                continue
            pt = r.font.size.pt if r.font.size else 18
            col = "#222"
            try:
                if r.font.color and r.font.color.type is not None:
                    col = "#" + str(r.font.color.rgb)
            except Exception:
                pass
            style = (f"font-size:{pt * scale / 72:.1f}px;color:{col};"
                     f"font-weight:{'700' if r.font.bold else '400'};"
                     f"font-family:{html.escape(r.font.name or 'Inter')},sans-serif")
            runs.append(f'<span style="{style}">{html.escape(r.text)}</span>')
        if runs:
            paras.append(f'<div style="text-align:{align};line-height:1.25;'
                         f'{"direction:rtl;" if rtl else ""}">{"".join(runs)}</div>')
    return "".join(paras)


def chart_summary(shape):
    try:
        ch = shape.chart
        cats = []
        try:
            cats = [str(c.label) for c in ch.plots[0].categories]
        except Exception:
            pass
        vals = []
        for s in ch.series:
            try:
                vals.append(f"{s.name}: " + ", ".join(
                    f"{v:g}" for v in s.values if v is not None))
            except Exception:
                pass
        bits = [f"📊 CHART ({str(ch.chart_type).split()[0]})"]
        if cats:
            bits.append("cats: " + ", ".join(cats[:8]))
        bits += vals[:3]
        return "<br>".join(html.escape(b) for b in bits)
    except Exception:
        return "📊 CHART"


def table_html(shape):
    rows = []
    for row in shape.table.rows:
        rows.append(" | ".join(html.escape(c.text) for c in row.cells))
    return "▦ TABLE<br>" + "<br>".join(rows[:6])


def main(argv):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("deck")
    ap.add_argument("--out", help="output HTML (default <deck>.contact.html)")
    ap.add_argument("--width", type=int, default=640, help="slide card width px")
    a = ap.parse_args(argv[1:])
    if not os.path.exists(a.deck):
        print(f"no such file: {a.deck}")
        return 2
    prs = Presentation(a.deck)
    W, H = prs.slide_width.inches, prs.slide_height.inches
    scale = a.width / W  # px per inch
    mm = morph_map(a.deck, len(prs.slides)) if hasattr(prs, "slides") else ""

    cards = []
    for idx, slide in enumerate(prs.slides, 1):
        shapes = list(walk(slide.shapes))
        bg = slide_bg(slide, shapes)
        els, words = [], 0
        for sh in shapes:
            if sh.left is None or not sh.width:
                continue
            l, t = sh.left.inches * scale, sh.top.inches * scale
            w, h = sh.width.inches * scale, sh.height.inches * scale
            base = (f"position:absolute;left:{l:.0f}px;top:{t:.0f}px;"
                    f"width:{w:.0f}px;height:{h:.0f}px;box-sizing:border-box;")
            if sh.shape_type == MSO_SHAPE_TYPE.CHART:
                els.append(f'<div style="{base}background:rgba(127,127,127,.14);'
                           f'border:1px dashed #888;padding:4px;font-size:10px;'
                           f'overflow:hidden">{chart_summary(sh)}</div>')
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                els.append(f'<div style="{base}background:rgba(127,127,127,.20);'
                           f'display:flex;align-items:center;justify-content:center;'
                           f'font-size:11px">🖼 IMAGE</div>')
            elif sh.has_table:
                els.append(f'<div style="{base}background:rgba(127,127,127,.12);'
                           f'border:1px solid #999;padding:4px;font-size:10px;'
                           f'overflow:hidden">{table_html(sh)}</div>')
            else:
                fill = solid(sh)
                inner = ""
                if sh.has_text_frame and sh.text_frame.text.strip():
                    inner = text_html(sh, scale)
                    words += len(sh.text_frame.text.split())
                els.append(
                    f'<div style="{base}'
                    f'{"background:" + fill + ";" if fill else ""}'
                    f'border-radius:{radius_for(sh)};overflow:hidden;'
                    f'padding:1px 3px">{inner}</div>')
        badge = {"M": "Morph", "t": "transition", ".": "no transition"}.get(
            mm[idx - 1] if idx - 1 < len(mm) else ".", "")
        notes = "notes ✓" if (slide.has_notes_slide and
                               slide.notes_slide.notes_text_frame.text.strip()) \
            else "no notes"
        cards.append(
            f'<figure><div class="slide" style="width:{a.width}px;'
            f'height:{H * scale:.0f}px;background:{bg}">{"".join(els)}</div>'
            f'<figcaption>#{idx} · {words} words · {badge} · {notes}</figcaption>'
            f'</figure>')

    doc = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<title>Contact sheet — {html.escape(os.path.basename(a.deck))}</title><style>
body{{background:#14161a;color:#e8eaed;font-family:Inter,system-ui,sans-serif;margin:24px}}
h1{{font-size:20px}} .sub{{color:#9aa0a6;font-size:13px;margin-bottom:18px}}
.grid{{display:flex;flex-wrap:wrap;gap:22px}}
figure{{margin:0}} .slide{{position:relative;border-radius:8px;overflow:hidden;
box-shadow:0 6px 24px rgba(0,0,0,.45)}}
figcaption{{font-size:12px;color:#9aa0a6;margin-top:6px}}
</style></head><body>
<h1>{html.escape(os.path.basename(a.deck))}</h1>
<div class="sub">{len(cards)} slides · {W:.2f}×{H:.2f} in · morph map <code>{mm}</code>
· wireframe proof sheet (tools/render.py) — geometry/colors/type are real; charts &amp; images are labeled blocks</div>
<div class="grid">{''.join(cards)}</div></body></html>
"""
    out = a.out or os.path.splitext(a.deck)[0] + ".contact.html"
    with open(out, "w", encoding="utf-8") as f:
        f.write(doc)
    print(f"contact sheet: {out} ({len(cards)} slides, {len(doc)//1024} KB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
