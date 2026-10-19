#!/usr/bin/env python3
"""Slide-library merge — cherry-pick slides from several decks into one.

Copies slide XML across files with full OPC relationship surgery: images,
native charts (incl. their embedded workbooks), media and hyperlink rels
travel with the slide; r:ids are remapped. Layout/master furniture from
the source does NOT transfer (shapes carry their own explicit styling in
generator decks; solid slide backgrounds are re-stamped). Notes copy too.

Usage:
  python tools/merge.py --out combined.pptx deckA.pptx:1,3-5 deckB.pptx:2
  python tools/merge.py --out c.pptx --morph talks/*.pptx   # whole decks
"""
import argparse
import copy
import os
import re
import sys

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(ROOT, "..", "generator"))
import generate_pptx as G  # noqa: E402

SKIP_RELTYPES = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE}


def unique_partname(dst_pkg, src_name):
    existing = {str(p.partname) for p in dst_pkg.iter_parts()}
    if src_name not in existing:
        return src_name
    m = re.match(r"^(.*?)(\d*)(\.[^.]+)$", src_name)
    stem, num, ext = m.group(1), int(m.group(2) or 0), m.group(3)
    while True:
        num += 1
        cand = f"{stem}{num}{ext}"
        if cand not in existing:
            return cand


def copy_part(src_part, dst_pkg, seen):
    """Deep-copy an OPC part (+ its internal rels graph) into dst package."""
    if src_part in seen:
        return seen[src_part]
    new = Part(PackURI(unique_partname(dst_pkg, str(src_part.partname))),
               src_part.content_type, dst_pkg, src_part.blob)
    seen[src_part] = new
    id_map = {}
    for rId, rel in list(src_part.rels.items()):
        if rel.is_external:
            nid = new.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            nid = new.relate_to(copy_part(rel.target_part, dst_pkg, seen),
                                rel.reltype)
        if nid != rId:
            id_map[rId] = nid
    if id_map and new._blob:  # keep in-XML references (e.g. chart externalData)
        blob = new._blob
        for old, nid in id_map.items():
            blob = blob.replace(f'"{old}"'.encode(), f'"{nid}"'.encode())
        new._blob = blob
    return new


def import_slide(src_slide, dst_prs, seen):
    blank = dst_prs.slide_layouts[6]
    new = dst_prs.slides.add_slide(blank)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    try:
        if src_slide.background.fill.type == MSO_FILL.SOLID:
            G.set_bg(new, str(src_slide.background.fill.fore_color.rgb))
    except Exception:
        pass
    dst_pkg = dst_prs.part.package
    id_map = {}
    for rId, rel in list(src_slide.part.rels.items()):
        if rel.reltype in SKIP_RELTYPES:
            continue
        if rel.is_external:
            nid = new.part.relate_to(rel.target_ref, rel.reltype,
                                     is_external=True)
        else:
            nid = new.part.relate_to(
                copy_part(rel.target_part, dst_pkg, seen), rel.reltype)
        if nid != rId:
            id_map[rId] = nid
    for shp in src_slide.shapes:
        new.shapes._spTree.insert_element_before(
            copy.deepcopy(shp._element), "p:extLst")
    if id_map:
        for el in new.shapes._spTree.iter():
            for k, v in list(el.attrib.items()):
                if v in id_map:
                    el.attrib[k] = id_map[v]
    if src_slide.has_notes_slide:
        txt = src_slide.notes_slide.notes_text_frame.text.strip()
        if txt:
            new.notes_slide.notes_text_frame.text = txt
    return new


def parse_selection(spec, count):
    if not spec:
        return list(range(1, count + 1))
    out = []
    for tok in spec.split(","):
        tok = tok.strip()
        if "-" in tok:
            a, b = tok.split("-", 1)
            out += range(int(a), int(b) + 1)
        elif tok:
            out.append(int(tok))
    return [n for n in out if 1 <= n <= count]


def main(argv):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("sources", nargs="+",
                    help="deck.pptx[:1,3-5] — no suffix = whole deck")
    ap.add_argument("--out", required=True)
    ap.add_argument("--morph", action="store_true",
                    help="re-stamp Morph on slides 2..N of the result")
    a = ap.parse_args(argv[1:])

    jobs = []
    for src in a.sources:
        path, _, sel = src.rpartition(":")
        if not path or not os.path.exists(path):
            if os.path.exists(src):
                path, sel = src, ""
            else:
                print(f"no such file: {src}")
                return 2
        jobs.append((path, sel))

    dst = None
    seen = {}
    total = 0
    for path, sel in jobs:
        src_prs = Presentation(path)
        if dst is None:
            dst = Presentation()
            dst.slide_width = src_prs.slide_width
            dst.slide_height = src_prs.slide_height
            # drop the default template's empty first slide? default has none
        picks = parse_selection(sel, len(src_prs.slides))
        print(f"{os.path.basename(path)}: slides {picks}")
        for n in picks:
            import_slide(src_prs.slides[n - 1], dst, seen)
            total += 1
    if not total:
        print("nothing selected")
        return 1
    if a.morph:
        for i, slide in enumerate(dst.slides):
            if i > 0:
                G.add_morph(slide, duration_ms=800)
    dst.save(a.out)
    print(f"merged {total} slides -> {a.out}"
          f"{' (Morph re-stamped)' if a.morph else ''}")
    print(f"verify: python tools/qa_deck.py {a.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
