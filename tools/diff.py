#!/usr/bin/env python3
"""Deck diff — changelog between two versions of a .pptx.

Compares slide count, per-slide text (added/removed lines), word counts,
Morph maps, notes coverage, palettes and fonts. Markdown out; exit 1 when
differences exist (diff(1) semantics — handy in scripts/CI), 0 when the
decks are textually identical.

Usage:
  python tools/diff.py old.pptx new.pptx [--out changelog.md]
"""
import argparse
import difflib
import os
import sys
from collections import Counter

from pptx import Presentation
from pptx.enum.dml import MSO_FILL

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from anatomy import morph_map, walk  # noqa: E402


def extract(path):
    prs = Presentation(path)
    slides, fills, fonts = [], Counter(), Counter()
    for slide in prs.slides:
        lines = []
        for sh in walk(slide.shapes):
            try:
                if sh.fill and sh.fill.type == MSO_FILL.SOLID:
                    fills[str(sh.fill.fore_color.rgb)] += 1
            except Exception:
                pass
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                txt = "".join(r.text for r in para.runs).strip()
                if txt:
                    lines.append(txt)
                for r in para.runs:
                    if r.font.name:
                        fonts[r.font.name] += 1
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append((lines, notes))
    return dict(n=len(slides), slides=slides, fills=fills, fonts=fonts,
                morph=morph_map(path, len(slides)),
                W=prs.slide_width.inches, H=prs.slide_height.inches)


def main(argv):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("old")
    ap.add_argument("new")
    ap.add_argument("--out", metavar="FILE")
    a = ap.parse_args(argv[1:])
    for p in (a.old, a.new):
        if not os.path.exists(p):
            print(f"no such file: {p}")
            return 2
    A, B = extract(a.old), extract(a.new)
    L = [f"# Deck diff — {os.path.basename(a.old)} → {os.path.basename(a.new)}",
         ""]
    changed = False

    if (A["W"], A["H"]) != (B["W"], B["H"]):
        L.append(f"- Canvas: {A['W']:.2f}×{A['H']:.2f} → {B['W']:.2f}×{B['H']:.2f} in")
        changed = True
    if A["n"] != B["n"]:
        L.append(f"- Slides: {A['n']} → {B['n']} ({B['n'] - A['n']:+d})")
        changed = True
    if A["morph"] != B["morph"]:
        L.append(f"- Morph map: `{A['morph']}` → `{B['morph']}`")
        changed = True

    fa, fb = set(A["fills"]), set(B["fills"])
    if fa != fb:
        changed = True
        if fb - fa:
            L.append("- Colors added: " + ", ".join(f"#{c}" for c in sorted(fb - fa)))
        if fa - fb:
            L.append("- Colors removed: " + ", ".join(f"#{c}" for c in sorted(fa - fb)))
    ta, tb = set(A["fonts"]), set(B["fonts"])
    if ta != tb:
        changed = True
        L.append(f"- Fonts: {sorted(ta)} → {sorted(tb)}")

    L.append("")
    total = max(A["n"], B["n"])
    for i in range(total):
        la, na = A["slides"][i] if i < A["n"] else ([], "")
        lb, nb = B["slides"][i] if i < B["n"] else ([], "")
        if la == lb and na == nb:
            continue
        changed = True
        if i >= A["n"]:
            L.append(f"## Slide {i + 1} — ADDED ({len(' '.join(lb).split())} words)")
            L += [f"  + {x}" for x in lb[:6]]
            continue
        if i >= B["n"]:
            L.append(f"## Slide {i + 1} — REMOVED ({len(' '.join(la).split())} words)")
            L += [f"  - {x}" for x in la[:6]]
            continue
        wa, wb = len(" ".join(la).split()), len(" ".join(lb).split())
        L.append(f"## Slide {i + 1} — changed ({wa} → {wb} words"
                 + (", notes edited" if na != nb else "") + ")")
        for line in difflib.unified_diff(la, lb, lineterm="", n=0):
            if line.startswith(("---", "+++", "@@")):
                continue
            if line.startswith(("-", "+")):
                L.append(f"  {line[:120]}")
        L.append("")

    if not changed:
        L.append("No textual, structural, palette, or font differences. "
                 "(Geometry-only moves still count as identical here — "
                 "compare contact sheets via tools/render.py for those.)")
    out = "\n".join(L) + "\n"
    print(out)
    if a.out:
        with open(a.out, "w", encoding="utf-8") as f:
            f.write(out)
        print(f"[written to {a.out}]")
    return 1 if changed else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
