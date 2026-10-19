#!/usr/bin/env python3
"""Handout — the leave-behind document version of any deck.

Slides are for presenting; documents are for reading (Duarte's
slideument warning, reference/12). This produces a clean, printable
HTML handout: per slide — headline, on-slide text, the full speaker
script, and talk time — plus a contents table with a running clock.
Print to PDF from the browser for a proper leave-behind that doesn't
betray the deck's leanness.

Usage:  python tools/handout.py deck.pptx [--out handout.html] [--wpm 130]
"""
import argparse
import datetime
import html
import os
import sys
from zipfile import ZipFile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from coach import fmt, slide_facts  # noqa: E402
from anatomy import walk  # noqa: E402
from pptx import Presentation  # noqa: E402


def slide_lines(slide):
    lines = []
    for sh in walk(slide.shapes):
        if sh.has_chart:
            lines.append("[chart]")
            continue
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs).strip()
            if t:
                lines.append(t)
    return lines


def main(argv):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("deck")
    ap.add_argument("--out", help="output HTML (default <deck>.handout.html)")
    ap.add_argument("--wpm", type=float, default=130)
    a = ap.parse_args(argv[1:])
    if not os.path.exists(a.deck):
        print(f"no such file: {a.deck}")
        return 2
    prs = Presentation(a.deck)
    title = prs.core_properties.title or os.path.basename(a.deck)
    try:
        xml0 = ZipFile(a.deck).read("ppt/slides/slide1.xml").decode(
            "utf-8", "ignore")
        rtl = 'rtl="1"' in xml0
    except Exception:
        rtl = False

    toc, secs, clock = [], [], 0.0
    for i, slide in enumerate(prs.slides, 1):
        headline, on_words, notes, _ = slide_facts(slide)
        nw = len(notes.split())
        sec = nw / a.wpm * 60
        lines = [l for l in slide_lines(slide) if l != headline]
        toc.append(f'<tr><td>{i}</td><td>{html.escape(headline[:70])}</td>'
                   f'<td>{fmt(sec)}</td><td>{fmt(clock)}</td></tr>')
        body = "".join(f"<li>{html.escape(l)}</li>" for l in lines[:14])
        script = html.escape(notes) if notes else \
            "<i>(no speaker script recorded)</i>"
        secs.append(
            f'<section><h2><span class="n">{i:02d}</span> '
            f'{html.escape(headline)}</h2>'
            f'<p class="meta">≈ {fmt(sec)} of talk · at {fmt(clock)} · '
            f'{on_words} words on slide</p>'
            f'<div class="cols"><div><h3>On the slide</h3><ul>{body}</ul></div>'
            f'<div><h3>The script</h3><p class="script">{script}</p></div>'
            f'</div></section>')
        clock += sec

    doc = f"""<!doctype html><html lang="{'ar' if rtl else 'en'}"{' dir="rtl"' if rtl else ''}>
<head><meta charset="utf-8"><title>Handout — {html.escape(title)}</title><style>
body{{font-family:Georgia,Tajawal,serif;max-width:880px;margin:40px auto;padding:0 22px;color:#1a1d21;line-height:1.55}}
h1{{font-family:Montserrat,Cairo,system-ui,sans-serif;font-size:30px}}
h2{{font-family:Montserrat,Cairo,system-ui,sans-serif;font-size:21px;margin-bottom:2px}}
h2 .n{{color:#0a9a8a}} h3{{font-size:12px;letter-spacing:.12em;text-transform:uppercase;color:#777}}
.meta{{color:#888;font-size:13px;margin-top:0}}
.cols{{display:grid;grid-template-columns:1fr 1.4fr;gap:26px}}
.script{{background:#f4f6f5;border-inline-start:3px solid #0a9a8a;padding:10px 14px}}
section{{border-top:1px solid #ddd;padding:18px 0}}
table{{border-collapse:collapse}} td{{padding:4px 16px 4px 0;border-bottom:1px solid #eee;font-size:14px}}
ul{{margin:6px 0;padding-inline-start:18px}} li{{margin:3px 0}}
@media print{{ section{{break-inside:avoid}} body{{margin:10px}} }}
</style></head><body>
<h1>{html.escape(title)} — handout</h1>
<p class="meta">{len(toc)} slides · scripted runtime ≈ {fmt(clock)} at {a.wpm:.0f} wpm ·
generated {datetime.date.today().isoformat()} by tools/handout.py · print me (Ctrl/Cmd+P)</p>
<h2>Contents</h2><table>{''.join(toc)}</table>
{''.join(secs)}
</body></html>
"""
    out = a.out or os.path.splitext(a.deck)[0] + ".handout.html"
    with open(out, "w", encoding="utf-8") as f:
        f.write(doc)
    print(f"handout: {out} ({len(toc)} slides, ≈{fmt(clock)} scripted)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
